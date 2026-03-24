"""Excel utilities for generating formatted product export files."""
import os
import re
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Exact column order matching the user's CSV template
CSV_TEMPLATE_COLUMNS = [
    'Category',
    'Product Name',
    'Brand',
    'Manufacturer',
    'Sale Price',
    'Discount Base Price',
    'Stock',
    'Lead Time',
    'Detailed Description',
    'Main Image',
    'Search Keywords',
    'Quantity',
    'Volume',
    'Weight',
    'Adult Only',
    'Taxable',
    'Parallel Import',
    'Overseas Purchase',
    'SKU',
    'Model Number',
    'Barcode',
    'Additional Image 1',
    'Additional Image 2',
]

# Column width configuration
COLUMN_WIDTHS = {
    'Category': 15,
    'Product Name': 45,
    'Brand': 20,
    'Manufacturer': 20,
    'Sale Price': 12,
    'Discount Base Price': 18,
    'Stock': 10,
    'Lead Time': 12,
    'Detailed Description': 60,
    'Main Image': 50,
    'Search Keywords': 40,
    'Quantity': 10,
    'Volume': 12,
    'Weight': 12,
    'Adult Only': 10,
    'Taxable': 10,
    'Parallel Import': 12,
    'Overseas Purchase': 15,
    'SKU': 18,
    'Model Number': 18,
    'Barcode': 18,
    'Additional Image 1': 50,
    'Additional Image 2': 50,
}

# URL columns that should be styled as links
URL_COLUMNS = {'Main Image', 'Additional Image 1', 'Additional Image 2'}

def build_excel(products, keyword, base_url, outputs_dir):
    """Build formatted Excel workbook from scraped product data."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products"

    # Use exact template columns only
    all_keys = CSV_TEMPLATE_COLUMNS.copy()

    # Style definitions
    header_fill = PatternFill("solid", fgColor="1A1A2E")
    header_font = Font(bold=True, color="E8F4FD", size=11, name="Calibri")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Side(style='thin', color='CCCCCC')
    cell_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

    # Write header row
    ws.row_dimensions[1].height = 35
    for ci, k in enumerate(all_keys, 1):
        cell = ws.cell(row=1, column=ci, value=k)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = cell_border

    # Data row styles
    alt_fill = PatternFill("solid", fgColor="F0F4FF")
    norm_fill = PatternFill("solid", fgColor="FFFFFF")
    data_font = Font(size=10, name="Calibri")
    link_font = Font(size=10, name="Calibri", color="0563C1", underline="single")

    # Write data rows
    for ri, prod in enumerate(products, 2):
        fill = alt_fill if ri % 2 == 0 else norm_fill
        ws.row_dimensions[ri].height = 20
        for ci, k in enumerate(all_keys, 1):
            val = prod.get(k, '')
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.fill = fill
            cell.border = cell_border
            cell.alignment = Alignment(vertical="center")
            # Style URL columns as links
            is_url = k in URL_COLUMNS and str(val).startswith('http')
            cell.font = link_font if is_url else data_font

    # Set column widths
    for ci, k in enumerate(all_keys, 1):
        ws.column_dimensions[get_column_letter(ci)].width = COLUMN_WIDTHS.get(k, 22)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Create Summary sheet
    ws2 = wb.create_sheet("Summary")
    ws2.column_dimensions['A'].width = 28
    ws2.column_dimensions['B'].width = 38

    # Calculate statistics using correct field names
    summary_rows = [
        ("Scrape Summary", ""),
        ("Website", base_url),
        ("Keyword", keyword),
        ("Total Products", len(products)),
        ("Date & Time", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Fields Captured", len(all_keys)),
        ("", ""),
        ("With Sale Price", sum(1 for p in products if p.get('Sale Price'))),
        ("With Main Image", sum(1 for p in products if p.get('Main Image'))),
        ("With Brand", sum(1 for p in products if p.get('Brand'))),
        ("With Description", sum(1 for p in products if p.get('Detailed Description'))),
        ("With SKU", sum(1 for p in products if p.get('SKU'))),
    ]

    for ri, (label, value) in enumerate(summary_rows, 1):
        ws2.row_dimensions[ri].height = 22
        cell_a = ws2.cell(row=ri, column=1, value=label)
        cell_b = ws2.cell(row=ri, column=2, value=value)
        if ri == 1:
            cell_a.font = Font(bold=True, size=13, color="FFFFFF", name="Calibri")
            cell_a.fill = PatternFill("solid", fgColor="1A1A2E")
        else:
            cell_a.font = Font(bold=True, size=11, name="Calibri")
            cell_b.font = Font(size=11, name="Calibri")

    # Generate output filename
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_keyword = re.sub(r'[^\w]', '_', keyword)[:20]
    filepath = os.path.join(outputs_dir, f"scrape_{safe_keyword}_{timestamp}.xlsx")
    wb.save(filepath)
    return filepath
